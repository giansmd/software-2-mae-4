"""Word (.docx) report builder (python-docx)."""
from __future__ import annotations

from datetime import date
from io import BytesIO

from docx.shared import Inches
from docx import Document
from matplotlib import pyplot as plt

from app.schemas.simulation import SimulationResult
from app.services.prediction import METHOD_LABELS


def build_docx_report(result: SimulationResult) -> bytes:
    doc = Document()
    doc.add_heading("MAE-4 — Reporte de Simulación", level=0)
    doc.add_paragraph(f"Escenario: {result.scenario_name}")
    doc.add_paragraph(
        f"Método de predicción: {METHOD_LABELS.get(result.method, result.method)}"
    )
    doc.add_paragraph(f"Nivel de restauración hidrológica: {result.restoration_level:.0f}%")
    doc.add_paragraph(f"Proyección de subida del nivel del mar: {result.sea_level_rise.value}")
    doc.add_paragraph(f"Año horizonte: {result.horizon_year}")
    doc.add_paragraph(f"Retención de carbono base: {result.baseline_carbon:.1f} tC")
    doc.add_paragraph(
        f"Retención proyectada a {result.horizon_year}: {result.horizon_carbon:.1f} tC "
        f"({result.horizon_variation_pct:+.1f}% vs base)"
    )
    doc.add_paragraph(
        "Hipótesis de restauración ≥ +25%: " + ("CUMPLE" if result.meets_hypothesis else "NO CUMPLE")
    )
    doc.add_heading("Interpretación y trazabilidad", level=2)
    doc.add_paragraph(
        "La variación compara la retención proyectada al año horizonte con la retención base. "
        "El cumplimiento de la hipótesis se evalúa frente al umbral de +25 %. La figura muestra "
        "la trayectoria anual de carbono retenido; la tabla resume carbono, área y cobertura."
    )
    doc.add_paragraph(
        f"Fuente: Gemelo Digital de Humedales · Método: {METHOD_LABELS.get(result.method, result.method)} · "
        f"Fecha: {date.today().isoformat()}"
    )
    if result.ml_horizon_variation_pct is not None:
        doc.add_paragraph(
            f"Comparación al horizonte — motor: {result.engine_horizon_variation_pct:+.1f}% · "
            f"modelo IA: {result.ml_horizon_variation_pct:+.1f}%"
        )

    doc.add_paragraph("")
    doc.add_heading("Serie temporal (resumen cada 5 años)", level=2)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    header = table.rows[0].cells
    for i, h in enumerate(["Año", "Retenido (tC)", "Área (ha)", "Cobertura", "Var %"]):
        header[i].text = h
    for p in result.series:
        if (p.year - result.series[0].year) % 5 == 0:
            cells = table.add_row().cells
            cells[0].text = str(p.year)
            cells[1].text = f"{p.retained_carbon:.1f}"
            cells[2].text = f"{p.effective_area_ha:.1f}"
            cells[3].text = f"{p.vegetation_cover:.2f}"
            cells[4].text = f"{p.variation_pct:.1f}"

    doc.add_paragraph(
        "Nota de tabla: carbono retenido en tC; área efectiva en ha; cobertura como proporción [0–1]; "
        "variación porcentual respecto al valor base. Se muestran observaciones cada cinco años."
    )
    doc.add_heading("Retención de carbono azul por año", level=2)
    fig, ax = plt.subplots(figsize=(8, 4.5))
    ax.plot([p.year for p in result.series], [p.retained_carbon for p in result.series], color="#0f766e", linewidth=2)
    ax.set_title("Proyección anual de carbono retenido")
    ax.set_xlabel("Año")
    ax.set_ylabel("Carbono retenido (tC)")
    ax.grid(alpha=0.25)
    fig.text(
        0.5, 0.02,
        f"Fuente: Gemelo Digital de Humedales · Método: {METHOD_LABELS.get(result.method, result.method)} · Fecha: {date.today().isoformat()}",
        ha="center", fontsize=7, color="#4b5563",
    )
    fig.tight_layout(rect=(0, 0.08, 1, 1))
    image = BytesIO()
    fig.savefig(image, format="png", dpi=600, bbox_inches="tight", pad_inches=0.1)
    plt.close(fig)
    image.seek(0)
    doc.add_picture(image, width=Inches(6.2))

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
