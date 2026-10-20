"""Excel (.xlsx) report builder (openpyxl) — full raw time series."""
from __future__ import annotations

from io import BytesIO
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference

from app.schemas.simulation import SimulationResult
from app.services.prediction import METHOD_LABELS


def build_xlsx_report(result: SimulationResult) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Serie_temporal"
    ws.append(["year", "retained_carbon_tc", "effective_area_ha", "vegetation_cover", "variation_pct"])
    for p in result.series:
        ws.append([p.year, p.retained_carbon, p.effective_area_ha, p.vegetation_cover, p.variation_pct])

    chart = LineChart()
    chart.title = "Retención de carbono azul por año"
    chart.y_axis.title = "Carbono retenido (tC)"
    chart.x_axis.title = "Año"
    chart.style = 13
    chart.height = 8
    chart.width = 16
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=ws.max_row), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
    chart.legend = None
    ws.add_chart(chart, "G2")

    summary = wb.create_sheet("Resumen")
    rows = [
        ("Campo", "Valor"),
        ("scenario", result.scenario_name),
        ("scenario_id", result.scenario_id),
        ("method", result.method),
        ("method_label", METHOD_LABELS.get(result.method, result.method)),
        ("restoration_level_pct", result.restoration_level),
        ("sea_level_rise", result.sea_level_rise.value),
        ("horizon_year", result.horizon_year),
        ("baseline_area_ha", result.baseline_area_ha),
        ("baseline_carbon_density", result.baseline_carbon_density),
        ("baseline_carbon_tc", result.baseline_carbon),
        ("horizon_carbon_tc", result.horizon_carbon),
        ("horizon_variation_pct", result.horizon_variation_pct),
        ("meets_hypothesis", result.meets_hypothesis),
        ("engine_horizon_variation_pct", result.engine_horizon_variation_pct),
        ("engine_horizon_carbon_tc", result.engine_horizon_carbon),
        ("ml_horizon_variation_pct", result.ml_horizon_variation_pct),
        ("ml_horizon_carbon_tc", result.ml_horizon_carbon),
        ("source", "Gemelo Digital de Humedales"),
        ("interpretation", "La variación compara carbono proyectado al horizonte con el valor base."),
        ("hypothesis_interpretation", "Cumple si la variación alcanza el umbral de +25 %."),
        ("series_note", "Carbono=tC; área=ha; cobertura=[0,1]; variación=% respecto al valor base."),
        ("generated_at", date.today().isoformat()),
    ]
    for r in rows:
        summary.append(list(r))

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
